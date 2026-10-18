from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'docs'/'Review_2_Report.docx'
EVIDENCE=ROOT/'docs'/'evidence'
NAVY='183B58'; TEAL='327C79'; GRAY='5D6B78'
d=Document(); sec=d.sections[0]
sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.78); sec.right_margin=Inches(.78)
for name,size,color in [('Normal',9.5,'263B4C'),('Title',28,NAVY),('Heading 1',17,NAVY),('Heading 2',12,TEAL),('Heading 3',10,NAVY)]:
 s=d.styles[name]; s.font.name='Aptos'; s.font.size=Pt(size); s.font.color.rgb=RGBColor.from_string(color)
 if name!='Normal': s.font.bold=True; s.paragraph_format.keep_with_next=True
 s.paragraph_format.space_after=Pt(5)
header=sec.header.paragraphs[0]; header.text='BCSE406L  |  NO SQL DATABASES  |  REVIEW 2'; header.runs[0].font.color.rgb=RGBColor.from_string(TEAL); header.runs[0].font.size=Pt(8)
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
rr=footer.add_run('Relational-to-NoSQL Schema Migration Tool  ·  '); rr.font.size=Pt(8); rr.font.color.rgb=RGBColor.from_string(GRAY)
fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

def add_table(head, rows, widths=None, size=7.8):
 t=d.add_table(rows=1,cols=len(head)); t.style='Light Shading Accent 1'; t.alignment=WD_TABLE_ALIGNMENT.CENTER
 hdr=OxmlElement('w:tblHeader'); hdr.set(qn('w:val'),'true'); t.rows[0]._tr.get_or_add_trPr().append(hdr)
 for i,h in enumerate(head):
  c=t.rows[0].cells[i]; c.text=h; sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),NAVY); c._tc.get_or_add_tcPr().append(sh)
  for r in c.paragraphs[0].runs: r.font.bold=True; r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string('FFFFFF')
 for j,row in enumerate(rows):
  cells=t.add_row().cells
  for i,val in enumerate(row):
   cells[i].text=str(val)
   if j%2: sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'F3F6F8'); cells[i]._tc.get_or_add_tcPr().append(sh)
   for p in cells[i].paragraphs:
    p.paragraph_format.space_after=Pt(1)
    for r in p.runs: r.font.size=Pt(size)
 if widths:
  for row in t.rows:
   for i,w in enumerate(widths): row.cells[i].width=Inches(w)
 d.add_paragraph().paragraph_format.space_after=Pt(1)
 return t

def p(text): d.add_paragraph(text)
def bullets(items):
 for text in items:
  x=d.add_paragraph(style='List Bullet'); x.paragraph_format.space_after=Pt(2); x.add_run(text)
def figure(name,caption):
 path=EVIDENCE/name
 if path.exists():
  x=d.add_paragraph(); x.alignment=WD_ALIGN_PARAGRAPH.CENTER; x.add_run().add_picture(str(path),width=Inches(6.45))
 else:
  raise FileNotFoundError(f'Required evidence screenshot missing: {path}')
 x=d.add_paragraph(caption,style='Caption'); x.alignment=WD_ALIGN_PARAGRAPH.CENTER

# Cover
x=d.add_paragraph(); x.paragraph_format.space_before=Pt(65); r=x.add_run('REVIEW 2  ·  IMPLEMENTATION REPORT'); r.bold=True; r.font.color.rgb=RGBColor.from_string(TEAL); r.font.size=Pt(10)
d.add_paragraph('Relational-to-NoSQL Schema Migration Tool',style='Title')
x=d.add_paragraph('Automated Denormalized Schema Recommendation with Trade-off Explanation'); x.runs[0].font.size=Pt(14); x.runs[0].font.color.rgb=RGBColor.from_string(GRAY)
add_table(['COURSE','INSTITUTION','TEAM'],[['BCSE406L – NoSQL Databases','Vellore Institute of Technology, Vellore','Anish Jaiswal · Savar Gambhir · Sakthi']],[1.8,2.3,2.5],8.2)
p('Submission date: 25 September 2026')
p('This report documents the Review 2 prototype built from the Review 1 design. It accepts relational DDL and query workload, derives a graph and workload signals, recommends MongoDB embedding or referencing with explicit rationale, and prepares generated schema artifacts. Claims are limited to implemented and verified behavior.')
add_table(['IMPLEMENTATION SNAPSHOT','VERIFIED EVIDENCE'],[['React/Vite frontend, FastAPI/Pydantic API, SQLGlot analysis, PyMongo persistence','Production frontend build succeeded; seven tests passed, including MongoDB API integration.'],['MongoDB is the primary target','Docker MongoDB 8.0 connected; seed, indexes, migration persistence/retrieval, live collection inspection, Extended JSON dump/restore verified.']],[2.25,4.35],8)
d.add_page_break()

d.add_heading('Executive summary',0)
p('The prototype preserves the Review 1 research direction: parse a normalized relational schema, analyze representative SQL queries, apply workload-aware denormalization rules, generate a MongoDB-oriented schema, and explain each decision. The API defines MongoDB collections, indexes, CRUD routes, migration persistence, retrieval, override regeneration, and export. The React dashboard provides input, analysis, recommendations, schema export, and history views.')
p('The React production build succeeded and the backend suite passed with 7 tests (0 failures; 3 non-blocking deprecation warnings). Docker Desktop runs MongoDB 8.0 on port 27017. The seeded database contains 15 customers, 20 products, and 40 orders; each order embeds line items. Database-backed project/workload CRUD, migration persistence and retrieval, collection inspection, and seed indexes were exercised. A database dump was restored and re-checked. Actual app and database captures appear throughout this report. Performance benchmarking and live relational-data ETL remain outside Review 2.')

d.add_heading('1. Project scope and Review 1 alignment',1)
d.add_heading('1.1 Problem and research gap',2)
p('Normalized relational designs support referential integrity and reduce redundancy, but high-read application paths often join multiple tables. Migration to a document store requires choosing which relationships to embed, reference, or duplicate. Review 1 identified a gap in end-to-end systems that combine DDL ingestion, workload-guided schema decisions, and a human-readable rationale for each decision.')
d.add_heading('1.2 Objectives preserved',2)
bullets(['Parse entities, attributes, primary keys, foreign keys, and relationships into an intermediate graph.','Analyze query table accesses, joins/co-access, filter, sort, and group attributes.','Recommend embed/reference choices from workload and fan-out evidence.','Generate MongoDB schema and explain benefits, read/write costs, and duplication concerns.','Permit relationship overrides and regenerate schema and explanations.'])
d.add_heading('1.3 Review 2 target',2)
p('MongoDB is the implemented target because documents and nested arrays represent denormalized aggregates naturally, while ObjectId references and indexes support independent access paths. Cassandra remains a possible later extension and is not claimed as implemented.')

d.add_heading('2. Architecture and implementation',1)
add_table(['Layer','Implemented responsibility','Technology'],[['Frontend','Dashboard, DDL/workload input, analysis, graph, recommendations, JSON export, history, Swagger link','React, Vite'],['API','Validation, health, CRUD, migration orchestration, override regeneration, result retrieval','FastAPI, Pydantic'],['Parsing','DDL tables, columns, PK/FK; SQL table aliases, joins, WHERE/ORDER/GROUP signals','SQLGlot'],['Decision pipeline','Graph, cardinality estimates, explicit rules, cost indicators, explanations, schema generator','Python'],['Persistence','Projects, schemas, workloads, runs, generated schemas, trade-off reports and indexes','MongoDB, PyMongo'],['Operations','Compose, seed script, sample SQL, tests, setup and backup instructions','Docker Compose, Pytest']],[1.0,3.55,2.05],7.6)
p('The data flow follows Review 1: SQL DDL + workload -> parser/analyzer -> unified relational graph -> cardinality profiler -> denormalization rules and cost indicators -> trade-off explainer -> MongoDB schema and persisted artifacts.')
add_table(['Input','Analysis','Decision and output','Persistence'],[['SQL DDL + query workload + row estimates','SQLGlot parser; workload extraction; FK graph; fan-out profile','Embedding/reference rules; indexes; schema; cost indicators and rationale','FastAPI/PyMongo writes run, source artifacts, generated schema and trade-off report to MongoDB']],[1.55,1.8,2.0,1.85],7.6)

d.add_heading('3. MongoDB database design',1)
d.add_heading('3.1 Collections and document structure',2)
add_table(['Collection','Purpose','Important fields and strategy'],[['projects','Project metadata','name, description, timestamps; independent document'],['schemas','Source DDL and parsed result','project_id, ddl, parsed graph; retained for traceability'],['workloads','Queries and extracted workload','project_id, queries, analysis'],['migration_runs','Complete input and generated analysis','project_id, input, result, status; embeds result for one-read retrieval'],['generated_schemas','Generated target schema','migration_id unique, project_id, schema'],['tradeoff_reports','Decision explanations','migration_id unique, project_id, items']],[1.25,2.25,3.35],7.7)
p('ObjectId values identify MongoDB documents. Runs embed their complete result to support one-read retrieval; generated schema and trade-off reports also have dedicated collections for direct access. Metadata documents reference the project/run IDs they belong to.')
d.add_heading('3.2 Embedding, references, and derived data',2)
p('The order-management model embeds order items into bounded order documents when workload co-access supports retrieving them together. High-fan-out customers or independently accessed products remain references when duplication risk outweighs read locality. The `order_total` field illustrates precomputation: item updates create a consistency/recomputation obligation. The tool generates recommendations; it does not copy data from a live relational source.')
d.add_heading('3.3 Indexes',2)
add_table(['Index','Collection','Supported access and trade-off'],[['name ascending','projects','Project lookup; index storage and write maintenance'],['project_id + created_at descending','schemas, workloads, migration_runs','Recent artifacts per project; compound index storage/write cost'],['unique migration_id','generated_schemas, tradeoff_reports','Direct result lookup and uniqueness enforcement'],['customer_id + order_date descending','orders (sample seed)','Customer order history; additional storage/update work'],['status + order_date descending','orders (sample seed)','Status filter/date ordering; additional storage/update work']],[1.8,1.8,3.25],7.5)
p('The API creates metadata indexes on startup. The seed creates sample-domain indexes. The live inspection route returned actual collection counts and index names; order indexes include unique `order_id`, `(customer_id, order_date desc)`, and `(status, order_date desc)`.')
figure('mongodb_indexes.jpg','Figure 3.1. Live collection counts and sample-domain MongoDB indexes, including the 40-document orders collection.')
figure('mongodb_documents.jpg','Figure 3.2. Actual MongoDB order documents with embedded item arrays.')

d.add_page_break()
d.add_heading('4. Database implementation and dataset',1)
d.add_heading('4.1 Setup',2)
code='docker compose -f database/docker-compose.yml up -d\npython -m pip install -r backend/requirements.txt\n$env:PYTHONPATH="backend"; python -m app.seed\n$env:PYTHONPATH="backend"; python -m uvicorn app.main:app --reload --port 8000\ncd frontend; npm install; npm run dev'
x=d.add_paragraph(); sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),'F1F4F6'); x._p.get_or_add_pPr().append(sh); r=x.add_run(code); r.font.name='Consolas'; r.font.size=Pt(8)
p('The Compose definition uses MongoDB 8.0 with port 27017, a persistent named data volume, and first-run collection/index initialization. The FastAPI startup hook also creates metadata collections and indexes when connected.')
d.add_heading('4.2 Demonstration dataset',2)
add_table(['Entity','Seed target','NoSQL representation'],[['customers','15','Standalone customer documents'],['products','20','Standalone product documents'],['orders','40','Order documents with nested items'],['order_items','80 line items','Embedded in order items array; seed leaves standalone collection empty']],[1.5,1.55,3.8],8)
p('These counts were observed in MongoDB after running the deterministic seed. The SQL DDL and four representative queries are stored in `database/sample_data/`. The 80 order items are nested within order documents rather than stored in an independent domain collection.')
d.add_heading('4.3 CRUD and persistence',2)
p('The API implements full project, schema and workload CRUD; migration creation and retrieval; recommendation/trade-off and generated-schema retrieval; export; and persisted override regeneration. Integration tests performed project/schema/workload create, read, update and delete against MongoDB, then created and retrieved a migration and its generated artifacts.')
figure('swagger.jpg','Figure 4.1. FastAPI OpenAPI/Swagger UI with the documented database and migration routes.')
figure('api_response.jpg','Figure 4.2. Actual MongoDB Inspector API response for the seeded orders collection, including 40 documents, indexes and sample records.')

d.add_heading('5. Core module implementation',1)
add_table(['Module','Input and processing','Output / evidence'],[['DDL parser','SQLGlot CREATE TABLE, columns/types, inline/table PK, FK REFERENCES','Structured tables, columns and keys; tests pass'],['Graph builder','One node per table and FK-directed parent-child edge','Graph from parsed schema; rendered in UI'],['Query analyzer','Resolve aliases and count accesses, joins, co-access, filters, sorts and groups','Actual query-derived workload counters; tested'],['Cardinality profiler','Supplied child/parent row ratio; <=1.5, <=8, >8 bands; missing -> unknown','Per-edge fan-out estimate; not inferred from DDL alone'],['Denormalization rules','Reference high/unknown fan-out or independent child access; otherwise bounded co-access can embed; support override','EMBED/REFERENCE, evidence and reason; override test passes'],['Cost estimator','Count selected relationships and joins; estimated avoided logical fetches = min(embedded edges, join occurrences)','Comparative heuristic, not latency benchmark'],['Trade-off explainer','Compose rationale from relationship, workload, fan-out, decision and cost','Evidence, reason, benefit and trade-off per edge'],['Schema generator','Map SQL types; preserve entities; embed selected children; suggest indexes','MongoDB schema; pipeline tests pass'],['Persistence/API','Pydantic validation, PyMongo collections/indexes, CRUD and result regeneration','MongoDB integration suite and live data inspection pass']],[1.25,3.45,2.15],7.1)

d.add_heading('6. Backend API implementation',1)
p('FastAPI OpenAPI is served at `/openapi.json` and Swagger UI at `/docs`. Pydantic validates required SQL text and non-empty query lists. Invalid SQL/workload requests return 422, invalid ObjectIds 400, missing records 404, and unavailable MongoDB 503.')
add_table(['Method','Endpoint','Purpose'],[['POST/GET/PUT/DELETE','/api/projects[/{project_id}]','Project CRUD'],['POST/GET/PUT/DELETE','/api/schemas[/{schema_id}]','Schema CRUD'],['POST/GET/PUT/DELETE','/api/workloads[/{workload_id}]','Workload CRUD'],['POST','/api/migrations','Run pipeline and persist artifacts'],['POST','/api/migrations/{id}/generate','Apply overrides and regenerate stored output'],['GET','/api/migrations[/{id}]','List history or retrieve run'],['GET','/api/migrations/{id}/recommendations','Retrieve decisions'],['GET','/api/migrations/{id}/tradeoffs','Retrieve explanation report'],['GET','/api/generated-schemas/{id}','Retrieve generated schema'],['GET','/api/database/collections','Inspect live collections/counts/index names']],[1.65,2.9,2.3],7.2)
p('Complete schemas are generated in the live OpenAPI document. `docs/api_reference.md` includes request shape and endpoint purposes.')

d.add_heading('7. Working prototype workflow',1)
bullets(['Open the dashboard and confirm MongoDB connection.','Enter project name, source row-count estimates, SQL DDL, and representative SQL queries.','Run the migration; parser and analyzer derive tables, keys, workload signals, graph and fan-out.','Inspect embed/reference decisions, evidence, benefits, trade-offs, and generated MongoDB schema.','Export JSON or override a relationship and regenerate.','Retrieve the persisted run from history and inspect MongoDB documents and indexes.'])
p('The screenshots below were captured from the running frontend, Swagger/API response and MongoDB-backed app. The browser run submitted the sample DDL and workload, produced recommendations and explanations, applied a relationship override, and retrieved persisted migration history.')
figure('dashboard.jpg','Figure 7.1. Application dashboard with live MongoDB connection and persisted metrics.')
figure('input.jpg','Figure 7.2. SQL DDL, row-count estimates and representative workload submitted in the application.')
figure('analysis.jpg','Figure 7.3. Parsed schema, workload signals, relational graph and generated recommendation result.')
figure('tradeoffs.jpg','Figure 7.4. Expanded per-relationship trade-off evidence, benefits and costs.')
figure('override.jpg','Figure 7.5. Persisted result after applying an EMBED override and regenerating the target schema.')
figure('history.jpg','Figure 7.6. Retrieved migration history from the MongoDB-backed application at screenshot capture time.')

d.add_heading('8. CRUD implementation',1)
add_table(['Operation','Route family','MongoDB action','Verification'],[['Create','POST projects/schemas/workloads/migrations','insert_one; migration companion artifacts','Passed in DB-backed integration test'],['Read','GET projects/schemas/workloads/runs/schema/report','find_one/find/sort','Passed; persisted migration and schema retrieved'],['Update','PUT project/schema/workload; POST generate','find_one_and_update/update_one','Passed for project/schema/workload; override regeneration observed'],['Delete','DELETE project/schema/workload','delete_one','Passed for project/schema/workload']],[1.0,2.55,1.8,1.5],7.2)
p('Automated integration evidence covers project/schema/workload create, retrieve, update and delete, plus migration create and persisted-result retrieval. The screenshot in Figure 4.2 demonstrates an actual collection read with records and indexes returned.')

d.add_heading('9. Sample workload and target model',1)
p('The workload covers customer order history, order/item reads, product lookups, product aggregation, and status/date filtering. The analyzer parses actual input queries, counts query table usage and joined table pairs, and identifies WHERE, ORDER BY, and GROUP BY columns. These counters are inputs to the decision engine rather than hard-coded presentation statistics.')
p('The generated MongoDB schema follows the recommendations: bounded order-item children may appear as nested arrays, while high fan-out or independently accessed relationships may use references. Product/customer key choice depends on the submitted workload and row-count estimates. Generated field types, embedded shapes, and indexes provide a reviewable target design; automatic transfer of existing relational data is outside this prototype.')

d.add_heading('10. Testing and validation',1)
add_table(['ID','Test case','Expected','Actual','Status'],[['T-01','DDL and primary/foreign key parsing','Tables and keys represented','Parser tests passed','PASS'],['T-02','Join and WHERE workload parsing','Tables, joins and filters counted','Analyzer tests passed','PASS'],['T-03','Recommendation, override and schema generation','Decision changes and target shape exists','Pipeline tests passed','PASS'],['T-04','Invalid DDL','Clear parse failure','Rejected with expected ValueError','PASS'],['T-05','Health, OpenAPI and input validation','Docs and validation responses available','FastAPI TestClient passed','PASS'],['T-06','MongoDB project/schema/workload CRUD and migration retrieval','Create/read/update/delete and persisted output succeed','MongoDB integration test passed','PASS'],['T-07','Seed records, indexes and live sample read','Expected counts and indexes returned','15 customers, 20 products, 40 orders; order indexes observed','PASS'],['T-08','Database export and restore','All dumped records and indexes restored','9 collections dumped/restored and rechecked','PASS'],['T-09','Frontend production build','Optimized frontend assets generated','Vite production build succeeded','PASS']],[.45,1.75,1.65,1.65,.75],7.0)
p('Actual backend test command: `python -m pytest backend/tests -q` — 7 passed, 0 failed, 3 deprecation warnings. The Vite production build also succeeded. The MongoDB integration case verifies database-backed CRUD, persisted migration retrieval and a seeded collection/index read.')

d.add_heading('11. Performance evaluation',1)
p('No latency or throughput benchmark was performed. The cost estimator uses comparative counters derived from relationships and workload joins; it does not predict execution time.')
add_table(['Measurement','Status'],[['DDL parsing time','PENDING — TO BE MEASURED FOR REVIEW 3'],['Migration end-to-end time','PENDING — TO BE MEASURED FOR REVIEW 3'],['Indexed vs non-indexed lookup','PENDING — TO BE MEASURED FOR REVIEW 3'],['CRUD response time','PENDING — TO BE MEASURED FOR REVIEW 3']],[2.8,3.8],8)

d.add_heading('12. Review 1 to Review 2 progress',1)
add_table(['Review 1 component','Review 2 status','Evidence'],[['DDL parser and graph','Implemented','Parser/key tests; graph generated from FKs'],['Query analyzer','Implemented','Join/filter workload test'],['Cardinality profiler','Implemented','Ratio classification in pipeline'],['Denormalization engine','Implemented','Rules and override test'],['Cost estimator/explainer','Implemented','Generated comparative indicators and rationale'],['Target schema generator','Implemented','Generated schema test'],['MongoDB persistence','Implemented and verified','DB integration test, collection inspector, backup/restore'],['Frontend/API','Built and database-backed demonstration verified','Vite build, app screenshots, FastAPI integration'],['Cassandra target','Not implemented','Review 3 candidate']],[1.8,2.6,2.2],7.5)

d.add_heading('13. Current implementation status',1)
d.add_heading('Completed for Review 2 prototype',2)
bullets(['Modular SQL parser, query analyzer, graph, fan-out and transparent recommendation pipeline.','MongoDB-oriented generator, PyMongo persistence routes, bootstrap, Compose, seed script, live indexes and Extended JSON backup/restore.','React dashboard, FastAPI API, OpenAPI docs, schema export, override regeneration and persisted history.','README, sample DDL/workload, API reference, architecture, integration tests, demo guide and rubric map.','Seven backend tests and the frontend production build pass; MongoDB integration and persistence are verified.'])
d.add_heading('Completed verification for Review 2',2)
bullets(['Docker MongoDB 8.0 responds on port 27017; health endpoint reports connected.','Seed dataset, sample indexes, project/schema/workload CRUD and migration artifacts were verified against MongoDB.','Backup export and restore were exercised; the restored database retained collections, records and indexes.','Real frontend, Swagger, API and database inspector screenshots are embedded in this report.'])
d.add_heading('Planned for Review 3',2)
bullets(['Live source RDBMS data extraction/ETL, rather than schema recommendation alone.','Broader SQL dialect and semantic/type coverage; formal schema coverage validation.','Measured and calibrated costs, larger workload/schema evaluation, and possible Cassandra generator.'])

d.add_heading('14. Challenges and solutions',1)
add_table(['Challenge','Cause','Response and status'],[['SQLGlot constraint variants','Inline and named/table constraints have different AST forms','Handle both forms; test pass'],['SQL aliases','Query columns refer to aliases','Resolve aliases to base tables; test pass'],['Missing cardinality data','DDL carries no actual table sizes','Accept explicit estimates; unknown ratios conservatively reference'],['Heuristic scope','Real cost depends on sizes, update rates and access latency','Expose evidence, cost caveats and override'],['Local runtime orchestration','Frontend, API and MongoDB are separate processes','Compose MongoDB and documented local startup; connected health verified']],[1.45,2.25,2.9],7.5)

d.add_heading('15. Conclusion',1)
p('The Review 2 prototype implements and demonstrates the Review 1 workload-aware relational-to-document schema workflow. The React frontend and FastAPI API run with MongoDB 8.0; DDL and workload analysis generate an inspectable graph, fan-out, recommendations, trade-offs and MongoDB schema; a migration result persists and is retrievable. Database-backed CRUD, seeded collections and indexes, backup/restore, seven backend tests, production frontend build and real screenshots are verified. This remains a prototype: the deterministic decision and cost model is heuristic, automatic ETL from a live source RDBMS is not implemented, and performance comparisons remain future work for Review 3.')

d.add_heading('16. References',1)
refs=[
'Bansal, N., Sachdeva, S., and Awasthi, L. K. (2023). Query-based denormalization using hypergraph (QBDNH): a schema transformation model for migrating relational to NoSQL databases. Knowledge and Information Systems, 66, 681–722.',
'Vajk, T., Fehér, P., Fekete, K., and Charaf, H. (2013). Denormalizing data into schema-free databases. CogInfoCom, 747–752.',
'Namdeo, B., and Suman, U. (2020). Schema design advisor model for RDBMS to NoSQL database migration. International Journal of Information Technology, 13, 277–286.',
'Bansal, N., Sachdeva, S., and Awasthi, L. K. (2023). Schema generation for document stores using a workload-driven approach. The Journal of Supercomputing, 80(3), 4000–4048.',
'Chen, L., Davoudian, A., and Liu, M. (2022). A workload-driven method for designing aggregate-oriented NoSQL databases. Data & Knowledge Engineering, 142, 102089.',
'Mior, M. J., Salem, K., Aboulnaga, A., and Liu, R. (2017). NoSE: Schema Design for NoSQL Applications. IEEE Transactions on Knowledge and Data Engineering, 29(10), 2275–2289.',
'Kim, H.-J., Ko, E.-J., Jeon, Y.-H., and Lee, K.-H. (2018). Techniques and guidelines for effective migration from RDBMS to NoSQL. The Journal of Supercomputing, 76(10), 7936–7950.',
'Dai, J. (2019). SQL to NoSQL: What to do and How. IOP Conference Series: Earth and Environmental Science, 234, 012080.',
'Alotaibi, O., and Pardede, E. (2019). Transformation of Schema from Relational Database (RDB) to NoSQL Databases. Data, 4(4), 148.',
'Kuszera, E. M., Peres, L. M., and Didonet Del Fabro, M. (2021). Exploring data structure alternatives in the RDB to NoSQL document store conversion process. Information Systems, 105, 101941.',
'Mahmood, A. A. (2018). Automated Algorithm for Data Migration from Relational to NoSQL Databases. Al-Nahrain Journal for Engineering Sciences, 21(1), 60–65.',
'Hillenbrand, A., Störl, U., Nabiyev, S., and Klettke, M. (2021). Self-adapting data migration in the context of schema evolution in NoSQL databases. Distributed and Parallel Databases, 40, 5–25.'
]
for ref in refs: d.add_paragraph(ref,style='List Number')
p('Bibliographic entries are reused from the supplied Review 1 and are not independently verified in this report.')

d.add_heading('Appendix A. Reproducible 7–10 minute demo',1)
bullets(['30 s: problem and Review 1 research gap.','45 s: architecture and MongoDB-only target.','45 s: MongoDB collections, sample documents and indexes.','60 s: submit sample DDL, row counts and workload.','90 s: explain graph, query patterns and fan-out.','90 s: show decisions, evidence, cost caveats and override regeneration.','45 s: export generated schema and retrieve persisted history.','45 s: Swagger CRUD, actual tests, remaining Review 3 work.'])

d.add_heading('Appendix B. Viva questions and answers',1)
qa=[('Why NoSQL?','The migration target is a workload-aligned, denormalized model; relational input remains supported.'),('Why MongoDB?','Nested documents fit aggregate reads; references and indexes support independent access.'),('Why not simply retain an RDBMS?','The project addresses workloads where repeated joins motivate evaluating denormalized reads; it does not claim RDBMS is universally worse.'),('Why embed?','Bounded fan-out plus frequent co-access can remove a separate fetch.'),('Why reference?','High/unbounded fan-out, independent access or independently changing records make duplication costly.'),('What is denormalization?','Nesting or duplicating related data for selected reads, trading off storage and consistency.'),('How is fan-out determined?','Child-row estimate divided by parent-row estimate; it is not measured unless counts are observed.'),('How does workload affect design?','Parsed table accesses, joins/co-access, filters and sorts feed the rules and index suggestions.'),('Embedding trade-off?','Fewer reads versus larger parents and update propagation.'),('Reference trade-off?','Less duplication versus extra relationship fetches.'),('Why indexes?','They support observed lookups, filters and sort paths; they consume space and add write work.'),('What happens during updates?','Embedded data updates the parent; precomputed values such as totals require recomputation.'),('How does the SQL parser work?','SQLGlot parses supported CREATE TABLE definitions and exposes columns and key constraints.'),('How is the graph built?','Tables are nodes and foreign keys form directed parent-child edges.'),('How does recommendation work?','Transparent rules use fan-out, co-access and independent child access; users may override.'),('How is schema validated?','Current tests verify output shape and key fidelity cases; formal semantic verification remains future work.'),('What about high cardinality?','The prototype favors references to bound document growth; thresholds are heuristic.'),('What are limitations?','Estimated counts, uncalibrated costs, SQL subset and missing update-rate/document-size modelling.'),('Why no Cassandra now?','Review 2 concentrates verification on MongoDB rather than claiming shallow support for two targets.'),('What is Review 3?','Live ETL, broader validation, calibrated benchmarks, and potentially Cassandra generation.')]
for q,a in qa:
 x=d.add_paragraph(); x.paragraph_format.space_after=Pt(4); r=x.add_run(q+' '); r.bold=True; x.add_run(a)

d.add_heading('Appendix C. Rubric-to-evidence mapping',1)
add_table(['Criterion','Project evidence','Section / artifact'],[['Database implementation','MongoDB collections, docs, nested items and indexes','3-4; Figures 3.1-3.2; seed and backup export'],['CRUD','POST/GET/PUT/DELETE and persisted state','6, 8; integration tests; Figures 4.1-4.2'],['Data modelling quality','Embed/reference, fan-out, document structure, indexes','3, 5, 9; Figures 7.3-7.5'],['Functional prototype','DDL/workload through persisted output','7; Figures 7.1-7.6'],['Progress/documentation','Git repository, README, OpenAPI, tests, dump/restore and demo','10-14; source, manifest and project docs']],[1.55,3.5,1.55],7.5)
d.add_heading('Appendix D. Evidence capture status',1)
p('Figures 3.1–3.2, 4.1–4.2 and 7.1–7.6 are genuine screenshots captured from the running local frontend, Swagger/API response and live MongoDB inspection. They are stored as original JPEG files in `docs/evidence/`. Backend test results and the backup manifest/log are reported as textual verification, not represented as screenshots.')

d.save(OUT)
print(OUT)
